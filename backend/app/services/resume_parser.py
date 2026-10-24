
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

import fitz  # PyMuPDF
from docx import Document


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract selectable text from a PDF."""
    pages = []

    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        if pdf.needs_pass:
            raise ValueError("Password-protected PDFs are not supported.")

        for page in pdf:
            text = page.get_text("text")
            if text.strip():
                page_text = text.strip()
                project_links = []
                blocks = page.get_text("blocks")
                project_heading = next(
                    (
                        block
                        for block in blocks
                        if block[4].strip().casefold() == "projects"
                    ),
                    None,
                )
                project_section_endings = {
                    "achievements",
                    "areas of interest",
                    "certifications",
                    "education",
                    "experience",
                    "professional summary",
                    "skills",
                    "technical skills",
                }
                project_section_end = next(
                    (
                        block[1]
                        for block in blocks
                        if project_heading
                        and block[1] > project_heading[3]
                        and block[4].strip().casefold() in project_section_endings
                    ),
                    float("inf"),
                )

                for link in page.get_links():
                    uri = link.get("uri", "")
                    link_rect = link.get("from")
                    if (
                        not project_heading
                        or not link_rect
                        or not project_heading[1] <= link_rect.y0 < project_section_end
                        or not uri.startswith(("https://", "http://"))
                    ):
                        continue

                    parsed_uri = urlparse(uri)
                    github_path = parsed_uri.path.strip("/").split("/")
                    if (
                        parsed_uri.netloc.casefold() in {"github.com", "www.github.com"}
                        and len(github_path) < 2
                    ):
                        continue

                    preceding_blocks = [
                        block
                        for block in blocks
                        if block[3] <= link_rect.y0 and block[4].strip()
                    ]
                    if not preceding_blocks:
                        continue

                    preceding_block = max(preceding_blocks, key=lambda block: block[3])
                    project_name = preceding_block[4].strip().splitlines()[0]
                    link_label = page.get_textbox(link_rect).strip() or "Link"
                    project_links.append(
                        f"Project link for {project_name} ({link_label}): {uri}"
                    )

                if project_links:
                    page_text += "\n\n" + "\n".join(project_links)
                pages.append(page_text)

    return "\n\n".join(pages)


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract text from DOCX paragraphs and tables."""
    document = Document(BytesIO(file_bytes))
    sections = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]
            if any(cells):
                sections.append(" | ".join(cells))

    return "\n".join(sections)


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Choose the parser based on the file extension."""
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(file_bytes)

    if extension == ".docx":
        return extract_docx_text(file_bytes)

    raise ValueError("Only PDF and DOCX files are supported.")
